"""Запуск: python run_checks.py (без сервера и API-ключа)."""
import json,os,tempfile
from pathlib import Path
from rag import build_index,retrieve,generate
from scheduler import schedule
from export_utils import excel,pdf,json_bytes
from gantt_utils import build_gantt
from auth_core import authenticate,allowed
from openpyxl import load_workbook
from io import BytesIO
D=json.loads(Path("data.json").read_text(encoding="utf-8"))
v,m=build_index(D["articles"])
assert retrieve(D["questions"][0],D["articles"],v,m)
os.environ.pop("LLM_API_KEY",None)
assert generate("проверка",retrieve(D["questions"][0],D["articles"],v,m))["answer_source"]=="template"
s=schedule(D["tasks"],D["workers"]);assert s["status"] in ("Optimal","Feasible")
x=excel(s);assert set(load_workbook(BytesIO(x)).sheetnames)=={"Расписание","Сетка","Загрузка","Не назначенные"}
p=pdf(s);assert p[:4]==b"%PDF"
assert json.loads(json_bytes(s))["status"]==s["status"]
assert len(build_gantt(s).data)==4
assert authenticate("admin","admin123")["role"]=="admin"
assert not allowed("student","create")
print("OK: retrieval, fallback, MOLSP, Excel, PDF, JSON, Gantt, roles")
